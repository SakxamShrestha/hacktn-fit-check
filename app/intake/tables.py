"""Read a clinic's equipment wish-list out of an uploaded spreadsheet.

Nothing here judges the request. Rows become RequestLine objects and
app.domain.marketplace does the validating, so the normalising rules live in
one place. An upload is treated as untrusted input: the columns are guessed
rather than demanded, the row count is capped, and every failure comes back
as a single IntakeError carrying a sentence the clinic can act on.
"""

from __future__ import annotations

import csv
import logging
import os
from collections.abc import Sequence
from dataclasses import dataclass
from io import BytesIO, StringIO
from itertools import islice

from openpyxl import load_workbook

from app.domain.marketplace import DEFAULT_LINE_QUANTITY, MIN_LINE_QUANTITY, RequestLine

LOGGER = logging.getLogger(__name__)

# One clinic, one order. Past this it is a dump of somebody's asset register
# and we would be filing hundreds of lines nobody will read.
MAX_ROWS = 200

DEVICE_HEADERS = frozenset({"device", "item", "equipment", "name", "description"})
QUANTITY_HEADERS = frozenset({"quantity", "qty", "count", "units", "number"})

CSV_SUFFIX = ".csv"
XLSX_SUFFIX = ".xlsx"
SUPPORTED_SUFFIXES = (CSV_SUFFIX, XLSX_SUFFIX)

# With no header to go on: the first column names the kind, the second counts
# it. That is the shape every hand-written list arrives in.
FALLBACK_DEVICE_COLUMN = 0
FALLBACK_QUANTITY_COLUMN = 1

Row = tuple[str, ...]


class IntakeError(ValueError):
    """An upload we could not turn into request lines, said in plain words."""


@dataclass(frozen=True)
class _Layout:
    """Which column holds what, and where the equipment rows start."""

    device: int
    quantity: int | None
    first_data_row: int


def parse_table(raw: bytes, filename: str) -> tuple[RequestLine, ...]:
    """Turn an uploaded .csv or .xlsx wish-list into request lines.

    Raises IntakeError when the file is not a spreadsheet we read, cannot be
    opened, or holds no equipment at all.
    """
    suffix = _suffix(filename)
    if suffix not in SUPPORTED_SUFFIXES:
        raise IntakeError(
            f"{_name(filename)} is not a spreadsheet we can read. Save it as "
            f"{' or '.join(SUPPORTED_SUFFIXES)} and upload it again."
        )
    if not raw:
        raise IntakeError(f"{_name(filename)} was empty.")

    rows = _csv_rows(raw) if suffix == CSV_SUFFIX else _xlsx_rows(raw)
    lines = _lines(rows)
    if not lines:
        raise IntakeError(
            f"No equipment was found in {_name(filename)}. Give it a column "
            "headed Device and one headed Quantity, with one item per row."
        )
    return lines


def to_quantity(value: object) -> int:
    """How many units a cell asks for, defaulting when it does not say.

    A blank, a word or a nonsense count is a clinic being imprecise, not a
    reason to reject the whole upload, so anything unreadable asks for one.
    """
    try:
        number = int(float(str(value).strip()))
    except (TypeError, ValueError, OverflowError):
        return DEFAULT_LINE_QUANTITY
    return number if number >= MIN_LINE_QUANTITY else DEFAULT_LINE_QUANTITY


def _suffix(filename: str) -> str:
    return os.path.splitext(filename.strip())[1].lower()


def _name(filename: str) -> str:
    """The file as the clinic sees it, for the error message."""
    return os.path.basename(filename.strip()) or "that file"


def _csv_rows(raw: bytes) -> tuple[Row, ...]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise IntakeError(
            "That CSV is not plain UTF-8 text. Re-export it from your "
            "spreadsheet program as CSV UTF-8 and upload it again."
        ) from exc
    try:
        reader = csv.reader(StringIO(text))
        return tuple(_row(cells) for cells in islice(reader, MAX_ROWS + 1))
    except csv.Error as exc:
        raise IntakeError(f"That CSV could not be read: {exc}") from exc


def _xlsx_rows(raw: bytes) -> tuple[Row, ...]:
    try:
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - one message, whatever openpyxl hit
        raise IntakeError(
            "That .xlsx file could not be opened. Re-save it from your "
            "spreadsheet program, or upload it as a CSV instead."
        ) from exc
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        return tuple(_row(cells) for cells in islice(rows, MAX_ROWS + 1))
    except Exception as exc:  # noqa: BLE001
        raise IntakeError(
            f"The first sheet of that {XLSX_SUFFIX} upload could not be read: {exc}"
        ) from exc
    finally:
        workbook.close()


def _row(cells: Sequence[object]) -> Row:
    return tuple(_text(cell) for cell in cells)


def _text(cell: object) -> str:
    return "" if cell is None else str(cell).strip()


def _lines(rows: Sequence[Row]) -> tuple[RequestLine, ...]:
    layout = _layout(rows)
    found: list[RequestLine] = []
    for row in rows[layout.first_data_row :]:
        name = _cell(row, layout.device)
        if not name:
            continue
        found.append(
            RequestLine(device_type=name, quantity=to_quantity(_cell(row, layout.quantity)))
        )
        if len(found) == MAX_ROWS:
            LOGGER.warning("wish-list truncated to its first %d rows", MAX_ROWS)
            break
    return tuple(found)


def _layout(rows: Sequence[Row]) -> _Layout:
    """Read the header row, or decide there isn't one."""
    for index, row in enumerate(rows):
        if not any(row):
            continue
        headers = [_normalise(cell) for cell in row]
        device = _column(headers, DEVICE_HEADERS)
        if device is None:
            break  # The first row with content is equipment, not headers.
        return _Layout(
            device=device,
            quantity=_column(headers, QUANTITY_HEADERS),
            first_data_row=index + 1,
        )
    return _Layout(
        device=FALLBACK_DEVICE_COLUMN,
        quantity=FALLBACK_QUANTITY_COLUMN,
        first_data_row=0,
    )


def _column(headers: Sequence[str], names: frozenset[str]) -> int | None:
    for index, header in enumerate(headers):
        if header in names:
            return index
    return None


def _normalise(cell: str) -> str:
    """A header as written, reduced to something we can compare."""
    return " ".join(cell.split()).lower()


def _cell(row: Row, index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    return row[index]
